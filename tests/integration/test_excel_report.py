"""Workbook generation and spreadsheet-safety integration tests."""

from datetime import datetime
from pathlib import Path
from uuid import uuid4

from openpyxl import load_workbook

from linkedin_automation.domain.enums import OutreachStatus, RunStatus
from linkedin_automation.domain.models import (
    Lead,
    MatchEvidence,
    OutreachRecord,
    ResearchRun,
    utc_now,
)
from linkedin_automation.infrastructure.reporting import ExcelReportWriter, escape_cell


def test_escapes_formula_prefixes() -> None:
    assert escape_cell('=HYPERLINK("bad")') == '\'=HYPERLINK("bad")'
    assert escape_cell(" normal") == " normal"


def test_writes_deterministic_safe_workbook(tmp_path: Path) -> None:
    run = ResearchRun.create(
        search_name="test",
        schema_version=1,
        definition_json="{}",
        definition_hash="b" * 64,
        dry_run=False,
    )
    now = utc_now()
    lead = Lead(
        id=uuid4(),
        run_id=run.id,
        profile_url="https://www.linkedin.com/in/ada-example/",
        full_name="=Dangerous Name",
        headline="AI Engineer",
        current_title="AI Engineer",
        company="Example",
        location="India",
        industry="Software",
        evidence=MatchEvidence(("title:AI Engineer",)),
        source_search="test",
        first_observed_at=now,
        last_observed_at=now,
        company_size_min=1,
        company_size_max=10,
        self_employed=True,
        about="Independent quantitative trader",
        experience=("Founder at Example",),
        education=("Example University",),
        skills=("Options", "Python"),
        connections="500+",
        followers="1,200",
    )
    writer = ExcelReportWriter(tmp_path)
    outreach = OutreachRecord(
        lead_id=lead.id,
        run_id=run.id,
        firm_type="family office",
        funds_gbp=1_000_000,
        fund_evidence_url="https://example.com/evidence",
        score=85,
        status=OutreachStatus.APPROVED,
        connection_message="Let's connect",
        follow_up_message="Thanks for connecting",
    )
    first = writer.write(run, [lead], [outreach], "leads.xlsx")
    second = writer.write(run, [lead], [outreach], "leads.xlsx")
    assert first == second
    workbook = load_workbook(first)
    assert workbook.sheetnames == ["Leads", "Outreach Review", "Run Summary"]
    sheet = workbook["Leads"]
    assert sheet.freeze_panes == "A2"
    assert sheet["A2"].value == "'=Dangerous Name"
    assert sheet["O2"].hyperlink.target == lead.profile_url
    assert sheet["G2"].value == "1-10"
    assert sheet["H2"].value == "Yes"
    assert sheet["I2"].value == "Independent quantitative trader"
    assert isinstance(sheet["R2"].value, datetime)
    assert sheet["B2"].alignment.wrap_text is True
    assert sheet.row_dimensions[2].height == 60
    assert workbook["Outreach Review"]["F2"].value == 1_000_000
    assert workbook["Outreach Review"]["J2"].value == "approved"
    assert workbook["Outreach Review"]["K2"].value == "not_sent"
    assert workbook["Outreach Review"]["L2"].value == "not_approved"
    assert workbook["Outreach Review"]["M2"].value == "not_sent"
    assert workbook["Run Summary"]["B5"].value == RunStatus.PENDING.value
    workbook.close()
