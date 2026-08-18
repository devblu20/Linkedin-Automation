"""Deterministic and spreadsheet-safe OpenPyXL report writer."""

from __future__ import annotations

import os
import tempfile
from datetime import UTC
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from linkedin_automation.application.exceptions import ReportError
from linkedin_automation.domain.models import Lead, OutreachRecord, ResearchRun

LEAD_HEADERS = (
    "Full Name",
    "Headline",
    "Current Title",
    "Company",
    "Location",
    "Industry",
    "Company Size",
    "Self Employed",
    "About",
    "Experience",
    "Education",
    "Skills",
    "Connections",
    "Followers",
    "LinkedIn Profile URL",
    "Match Evidence",
    "Source Search",
    "Collected At (UTC)",
    "Notes",
)

OUTREACH_HEADERS = (
    "Full Name",
    "Company",
    "Current Title",
    "LinkedIn Profile URL",
    "Firm Type",
    "Verified Funds (GBP)",
    "Fund Evidence URL",
    "Fund Evidence Notes",
    "Qualification Score",
    "Review Status",
    "Connection Request Status",
    "Connection Approved",
    "Message Status",
    "Connection Message Draft",
    "Follow-up Message Draft",
    "Updated At (UTC)",
)


def escape_cell(value: str) -> str:
    """Prevent untrusted text from becoming an Excel formula."""
    if value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r", "\n")):
        return f"'{value}"
    return value


def _style_header(sheet: Worksheet) -> None:
    bottom_border = Border(bottom=Side(style="thin", color="9CA3AF"))
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(fill_type="solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
        cell.border = bottom_border
    sheet.row_dimensions[1].height = 24
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False


def _set_widths(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width


class ExcelReportWriter:
    """Write reports atomically inside an explicitly configured artifact root."""

    def __init__(self, artifact_dir: Path) -> None:
        self._artifact_dir = artifact_dir.resolve()

    def write(
        self,
        run: ResearchRun,
        leads: list[Lead],
        outreach: list[OutreachRecord],
        file_name: str,
    ) -> Path:
        target = (self._artifact_dir / file_name).resolve()
        if target.parent != self._artifact_dir or target.suffix.casefold() != ".xlsx":
            raise ReportError("report target must be an .xlsx file inside the artifact directory")
        self._artifact_dir.mkdir(parents=True, exist_ok=True)
        workbook = self._build_workbook(run, leads, outreach)
        temporary_path: Path | None = None
        try:
            with tempfile.NamedTemporaryFile(
                prefix=f".{target.stem}-", suffix=".xlsx", dir=self._artifact_dir, delete=False
            ) as temporary:
                temporary_path = Path(temporary.name)
            workbook.save(temporary_path)
            verified = load_workbook(temporary_path, read_only=True)
            verified.close()
            os.replace(temporary_path, target)
        except (OSError, ValueError) as error:
            if temporary_path is not None:
                temporary_path.unlink(missing_ok=True)
            raise ReportError(f"could not create Excel report: {target.name}") from error
        return target

    @staticmethod
    def _build_workbook(
        run: ResearchRun, leads: list[Lead], outreach: list[OutreachRecord]
    ) -> Workbook:
        workbook = Workbook()
        lead_sheet = workbook.active
        lead_sheet.title = "Leads"
        lead_sheet.append(LEAD_HEADERS)
        for lead in sorted(leads, key=lambda item: item.profile_url):
            lead_sheet.append(
                (
                    escape_cell(lead.full_name),
                    escape_cell(lead.headline),
                    escape_cell(lead.current_title),
                    escape_cell(lead.company),
                    escape_cell(lead.location),
                    escape_cell(lead.industry),
                    (
                        f"{lead.company_size_min or 1}-{lead.company_size_max}"
                        if lead.company_size_max is not None
                        else "Unknown"
                    ),
                    "Yes" if lead.self_employed is True else "No/Unknown",
                    escape_cell(lead.about),
                    escape_cell(" | ".join(lead.experience)),
                    escape_cell(" | ".join(lead.education)),
                    escape_cell(" | ".join(lead.skills)),
                    escape_cell(lead.connections),
                    escape_cell(lead.followers),
                    lead.profile_url,
                    escape_cell("; ".join(lead.evidence.matched)),
                    escape_cell(lead.source_search),
                    lead.last_observed_at.astimezone(UTC).replace(tzinfo=None),
                    escape_cell(lead.notes),
                )
            )
            url_cell = lead_sheet.cell(row=lead_sheet.max_row, column=15)
            url_cell.hyperlink = lead.profile_url
            url_cell.style = "Hyperlink"
        for row_number in range(2, lead_sheet.max_row + 1):
            lead_sheet.row_dimensions[row_number].height = 60
            for column_number in range(1, len(LEAD_HEADERS) + 1):
                cell = lead_sheet.cell(row=row_number, column=column_number)
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=column_number in {2, 3, 4, 9, 10, 11, 12, 16, 17, 19},
                )
            lead_sheet.cell(row=row_number, column=18).number_format = "yyyy-mm-dd hh:mm:ss"
        _style_header(lead_sheet)
        _set_widths(
            lead_sheet, (24, 50, 30, 30, 28, 24, 16, 15, 50, 60, 40, 40, 14, 14, 44, 36, 52, 22, 30)
        )

        outreach_sheet = workbook.create_sheet("Outreach Review")
        outreach_sheet.append(OUTREACH_HEADERS)
        leads_by_id = {lead.id: lead for lead in leads}
        for record in sorted(outreach, key=lambda item: (-item.score, str(item.lead_id))):
            outreach_lead = leads_by_id.get(record.lead_id)
            if outreach_lead is None:
                continue
            outreach_sheet.append(
                (
                    escape_cell(outreach_lead.full_name),
                    escape_cell(outreach_lead.company),
                    escape_cell(outreach_lead.current_title),
                    outreach_lead.profile_url,
                    escape_cell(record.firm_type),
                    record.funds_gbp,
                    record.fund_evidence_url,
                    escape_cell(record.fund_evidence_notes),
                    record.score,
                    record.status.value,
                    (
                        "sent"
                        if record.status.value
                        in {"connection_sent", "connected", "message_ready", "message_sent"}
                        else "not_sent"
                    ),
                    "approved"
                    if record.status.value in {"connected", "message_ready", "message_sent"}
                    else "not_approved",
                    "sent" if record.status.value == "message_sent" else "not_sent",
                    escape_cell(record.connection_message),
                    escape_cell(record.follow_up_message),
                    record.updated_at.astimezone(UTC).replace(tzinfo=None),
                )
            )
            for column in (4, 7):
                cell = outreach_sheet.cell(row=outreach_sheet.max_row, column=column)
                if cell.value:
                    cell.hyperlink = str(cell.value)
                    cell.style = "Hyperlink"
        for row_number in range(2, outreach_sheet.max_row + 1):
            outreach_sheet.row_dimensions[row_number].height = 72
            for column_number in range(1, len(OUTREACH_HEADERS) + 1):
                outreach_sheet.cell(row=row_number, column=column_number).alignment = Alignment(
                    vertical="top", wrap_text=column_number in {5, 8, 14, 15}
                )
            outreach_sheet.cell(row=row_number, column=6).number_format = "£#,##0"
            outreach_sheet.cell(row=row_number, column=16).number_format = "yyyy-mm-dd hh:mm:ss"
        _style_header(outreach_sheet)
        _set_widths(
            outreach_sheet,
            (24, 28, 28, 42, 22, 20, 42, 36, 18, 20, 22, 20, 18, 54, 54, 22),
        )

        summary = workbook.create_sheet("Run Summary")
        summary.append(("Field", "Value"))
        rows = (
            ("Run ID", str(run.id)),
            ("Search Name", escape_cell(run.search_name)),
            ("Definition Hash", run.definition_hash),
            ("Status", run.status.value),
            ("Started At (UTC)", run.created_at.astimezone(UTC).replace(tzinfo=None)),
            ("Updated At (UTC)", run.updated_at.astimezone(UTC).replace(tzinfo=None)),
            ("Observed", run.counters.observed),
            ("Accepted", run.counters.accepted),
            ("Rejected", run.counters.rejected),
            ("Duplicated", run.counters.duplicated),
            ("Failed", run.counters.failed),
            ("Checkpoint Page", run.checkpoint_page),
        )
        for row in rows:
            summary.append(row)
        _style_header(summary)
        _set_widths(summary, (28, 72))
        for row_number in range(2, summary.max_row + 1):
            summary.row_dimensions[row_number].height = 22
            summary.cell(row=row_number, column=1).font = Font(bold=True, color="374151")
            summary.cell(row=row_number, column=1).alignment = Alignment(vertical="center")
            summary.cell(row=row_number, column=2).alignment = Alignment(vertical="center")
        for row_number in (6, 7):
            summary.cell(row=row_number, column=2).number_format = "yyyy-mm-dd hh:mm:ss"
        for row_number in range(8, 14):
            summary.cell(row=row_number, column=2).number_format = "#,##0"
        workbook.properties.creator = "LinkedIn Lead Research Automation"
        workbook.properties.created = run.created_at.replace(tzinfo=None)
        workbook.properties.modified = run.updated_at.replace(tzinfo=None)
        return workbook
